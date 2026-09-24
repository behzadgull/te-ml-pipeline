"""
Build paper/templates/reference.docx, the pandoc reference document for the
journal-neutral submission manuscript (see scripts/build_paper.*).

Starts from pandoc's own default reference document
(`pandoc -o custom.docx --print-default-data-file reference.docx`) and edits
its styles: Times New Roman throughout; Normal and Body Text at 12 pt, 1.5
line spacing, justified; headings bold and black; captions at 10 pt; a plain
grid table style; A4 with 2.5 cm margins; page numbers in the footer; and
continuous line numbering that does not restart on each page.

Usage (repository root; needs python-docx):
    python paper/templates/make_reference_docx.py [path/to/pandoc]
"""

import copy
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path("paper/templates/reference.docx")
FONT = "Times New Roman"
A4_TWIPS = (11906, 16838)     # 21.0 x 29.7 cm
MARGIN_CM = 2.5


def _set_fonts(rpr, font=FONT):
    """Point every font slot of a run-properties element at `font`, dropping theme-font attributes."""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        del rfonts.attrib[attr]
    for slot in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{slot}"), font)


def _style(doc, style_id):
    for st in doc.styles:
        if st.style_id == style_id:
            return st
    raise KeyError(style_id)


def _format(st, size=None, bold=None, italic=None, spacing=None, align=None, before=None, after=None, keep_next=None):
    """Set font (Times New Roman, black) and paragraph formatting on a paragraph style."""
    _set_fonts(st.element.get_or_add_rPr())
    f = st.font
    f.name = FONT
    f.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    pf = st.paragraph_format
    if spacing is not None:
        pf.line_spacing = spacing
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if keep_next is not None:
        pf.keep_with_next = keep_next


def _add_field(paragraph, instr):
    """Append a complex field (e.g. PAGE) to a paragraph."""
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        run = paragraph.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            run.text = text
        run.font.name = FONT
        run.font.size = Pt(10)


def _grid_table_style(doc):
    """Replace pandoc's 'Table' style borders with a simple all-cells grid."""
    st = _style(doc, "Table")
    el = st.element
    for tag in ("w:tblPr", "w:tblStylePr"):
        pass
    tblpr = el.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = OxmlElement("w:tblPr")
        el.append(tblpr)
    for child in list(tblpr):
        if child.tag in (qn("w:tblBorders"), qn("w:tblCellMar")):
            tblpr.remove(child)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for edge, w in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        m = OxmlElement(f"w:{edge}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tblpr.append(mar)
    # drop conditional formatting that would add shading or extra rules, keep header-row bold
    for tsp in el.findall(qn("w:tblStylePr")):
        el.remove(tsp)
    tsp = OxmlElement("w:tblStylePr")
    tsp.set(qn("w:type"), "firstRow")
    rpr = OxmlElement("w:rPr")
    rpr.append(OxmlElement("w:b"))
    tsp.append(rpr)
    el.append(tsp)


def build(pandoc="pandoc"):
    with tempfile.TemporaryDirectory() as tmp:
        default = Path(tmp) / "custom.docx"
        subprocess.run([pandoc, "-o", str(default), "--print-default-data-file", "reference.docx"], check=True)
        doc = Document(str(default))

    # document defaults: Times New Roman, no theme fonts
    styles_el = doc.styles.element
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is not None:
        rpr = dd.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            _set_fonts(rpr)

    J, L, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER
    _format(_style(doc, "Normal"), size=12, spacing=1.5, align=J)
    for sid in ("BodyText", "FirstParagraph", "BlockText"):
        _format(_style(doc, sid), size=12, spacing=1.5, align=J, before=0, after=6)
    _format(_style(doc, "Compact"), size=10, spacing=1.0, align=L, before=0, after=0)   # table cells and lists
    _format(_style(doc, "Title"), size=16, bold=True, italic=False, spacing=1.0, align=C, before=0, after=12)
    _format(_style(doc, "Author"), size=12, italic=False, spacing=1.0, align=C, before=0, after=6)
    _format(_style(doc, "Date"), size=12, spacing=1.0, align=C)
    _format(_style(doc, "AbstractTitle"), size=12, bold=True, spacing=1.0, align=L)
    _format(_style(doc, "Abstract"), size=12, spacing=1.5, align=J)
    for sid, size in (("Heading1", 14), ("Heading2", 12), ("Heading3", 12), ("Heading4", 12)):
        _format(_style(doc, sid), size=size, bold=True, italic=False, spacing=1.0, align=L,
                before=14 if sid == "Heading1" else 10, after=6, keep_next=True)
    for sid in ("Caption", "ImageCaption", "TableCaption"):
        _format(_style(doc, sid), size=10, bold=False, italic=False, spacing=1.0, align=J, before=4, after=10,
                keep_next=(sid == "TableCaption"))     # a table caption stays with its table
    _format(_style(doc, "CaptionedFigure"), spacing=1.0, align=C, keep_next=True)   # a figure stays with its caption
    _format(_style(doc, "Figure"), spacing=1.0, align=C, keep_next=True)
    _format(_style(doc, "Bibliography"), size=10, spacing=1.15, align=L, before=0, after=4)
    _format(_style(doc, "FootnoteText"), size=10, spacing=1.0, align=J)
    _grid_table_style(doc)

    # section: A4, 2.5 cm margins, page number footer, continuous line numbers
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(MARGIN_CM))
    sec.header_distance = Cm(1.25)
    sec.footer_distance = Cm(1.25)
    sec.gutter = Cm(0)
    footer = sec.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = C
    _add_field(p, "PAGE")
    sectpr = sec._sectPr
    for old in sectpr.findall(qn("w:lnNumType")):
        sectpr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    pgmar = sectpr.find(qn("w:pgMar"))
    pgmar.addnext(ln)          # schema order: pgSz, pgMar, paperSrc, pgBorders, lnNumType, ...

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "pandoc")
